#!/usr/bin/env python3
"""Genera CeraLux_Creativos.xlsx: biblioteca de creativos (hook + cuerpo) con su rendimiento real.

    python creativos/generar_creativos.py

Lee biblioteca.json (anuncios, hooks, cuerpos, ajustes) e historico.json (Meta por anuncio y día + pedidos
reales de Shopify). Si el Excel ya existe, antes de regenerarlo recoge lo que se haya editado a mano
(hooks, cuerpos, notas, ajustes) y lo guarda en biblioteca.json, así no se pierde nada.
"""
import datetime as dt
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = Path(__file__).resolve().parent
BIB, HIST, SALIDA = AQUI / "biblioteca.json", AQUI / "historico.json", AQUI / "CeraLux_Creativos.xlsx"

FUENTE, TINTA, GRIS, AZUL = "Arial", "1B2537", "6B7686", "0000FF"
FONDO_INPUT, FONDO_SUB, FONDO_CARD = "FFF4CC", "E6EAF0", "F3F5F8"
VERDE, AMBAR, ROJO, VERDE_SUAVE, AMBAR_SUAVE, ROJO_SUAVE = "1E8449", "B9770E", "C0392B", "D5F0DD", "FBEBD0", "F8DEDA"
COLOR_ANGULO = ["DCE8F7", "F7E6D5", "E3F1E1", "EFE3F5", "F5F0D6", "DDEFF0"]
EUR = '#,##0.00 "€";[Red]-#,##0.00 "€";"-"'
EUR0 = '#,##0 "€";[Red]-#,##0 "€";"-"'
ENT = '#,##0;-#,##0;"-"'
PCT = '0.0%;-0.0%;"-"'
DEC = '0.00;-0.00;"-"'
FECHA = "dd/mm/yyyy"
H0, HN = 5, 5000          # filas de datos en Histórico
C0 = 5                    # primera fila de datos en Creativos / Cuerpos / Pedidos
HIS = "'Histórico'"
fino = Side(style="thin", color="D5DBE3")
BORDE = Border(left=fino, right=fino, top=fino, bottom=fino)


def font(bold=False, color="000000", size=10, italic=False):
    return Font(name=FUENTE, bold=bold, color=color, size=size, italic=italic)


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def celda(ws, ref, valor, *, bold=False, color="000000", size=10, fmt=None, fondo=None, align=None,
          wrap=False, italic=False, borde=None, valign="center"):
    c = ws[ref]
    c.value = valor
    c.font = font(bold, color, size, italic)
    if fmt:
        c.number_format = fmt
    if fondo:
        c.fill = fill(fondo)
    c.alignment = Alignment(horizontal=align, vertical=valign, wrap_text=wrap)
    if borde:
        c.border = borde
    return c


def entrada(ws, ref, valor, fmt=None, wrap=False, nota=None):
    c = celda(ws, ref, valor, color=AZUL, fmt=fmt, fondo=FONDO_INPUT, borde=BORDE, wrap=wrap, valign="top" if wrap else "center")
    if nota:
        c.comment = Comment(nota, "CeraLux")
    return c


def titulo(ws, texto, sub):
    celda(ws, "B1", texto, bold=True, color=TINTA, size=16)
    celda(ws, "B2", sub, color=GRIS, size=9, italic=True)


def seccion(ws, rango, texto):
    ini = rango.split(":")[0]
    ws.merge_cells(rango)
    celda(ws, ini, texto, bold=True, color="FFFFFF", fondo=TINTA)
    ws[ini].alignment = Alignment(horizontal="left", vertical="center", indent=1)


def cabecera(ws, fila, col0, textos, entradas=()):
    for i, t in enumerate(textos):
        c = ws.cell(row=fila, column=col0 + i, value=t)
        es_in = (col0 + i) in entradas
        c.font = font(True, "5A4300" if es_in else TINTA, 9)
        c.fill = fill(FONDO_INPUT if es_in else FONDO_SUB)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDE


def nombre(wb, n, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=ref)


def col(n):
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def S(letra, crit_col=None, crit=None):
    """Suma de una columna del Histórico en el periodo del Panel (y opcionalmente con un criterio)."""
    extra = f",{HIS}!${crit_col}${H0}:${crit_col}${HN},{crit}" if crit_col else ""
    return (f"SUMIFS({HIS}!${letra}${H0}:${letra}${HN}{extra},{HIS}!$A${H0}:$A${HN},\">=\"&DESDE,"
            f"{HIS}!$A${H0}:$A${HN},\"<=\"&HASTA)")


# ---------------------------------------------------------------- datos
bib = json.loads(BIB.read_text(encoding="utf-8"))
hist = json.loads(HIST.read_text(encoding="utf-8"))
cfg = bib["config"]

# 1) recoger lo editado a mano en el Excel anterior
if SALIDA.exists():
    old = load_workbook(SALIDA)
    if "Creativos" in old.sheetnames:
        ws = old["Creativos"]
        por_id = {a["id"]: a for a in bib["anuncios"]}
        for r in range(C0, ws.max_row + 1):
            aid = ws[f"A{r}"].value
            if aid and str(aid) in por_id:
                a = por_id[str(aid)]
                for k, letra in (("cuerpo", "E"), ("formato", "F"), ("tipo_hook", "G"), ("concepto", "H"), ("hook", "I"), ("notas", "L")):
                    v = ws[f"{letra}{r}"].value
                    a[k] = "" if v is None else str(v)
    if "Cuerpos" in old.sheetnames:
        ws = old["Cuerpos"]
        por_id = {c["id"]: c for c in bib["cuerpos"]}
        campos = (("nombre", "B"), ("angulo", "C"), ("formato", "D"), ("apertura", "E"), ("solucion", "F"), ("mecanismo", "G"),
                  ("prueba", "H"), ("oferta", "I"), ("cta", "J"), ("precio", "N"), ("notas", "O"))
        for r in range(C0, ws.max_row + 1):
            cid = ws[f"A{r}"].value
            if not cid:
                continue
            cid = str(cid).strip()
            c = por_id.get(cid)
            if c is None:
                c = {"id": cid}
                bib["cuerpos"].append(c)
                por_id[cid] = c
            for k, letra in campos:
                v = ws[f"{letra}{r}"].value
                c[k] = "" if v is None else str(v)
    if "Config" in old.sheetnames:
        ws = old["Config"]
        for r in range(4, 20):
            k, v = ws[f"F{r}"].value, ws[f"C{r}"].value
            if k in cfg and v is not None:
                cfg[k] = v.date().isoformat() if isinstance(v, dt.datetime) else v
    if "Pedidos" in old.sheetnames:
        ws = old["Pedidos"]
        for r in range(C0, ws.max_row + 1):
            f, v = ws[f"A{r}"].value, ws[f"B{r}"].value
            if isinstance(f, (dt.datetime, dt.date)) and v is not None:
                hist["pedidos_reales"].setdefault((f.date() if isinstance(f, dt.datetime) else f).isoformat(), v)

# 2) anuncios nuevos que aparecen en Meta y aún no están en la biblioteca
conocidos = {a["id"] for a in bib["anuncios"]}
for fila in hist["filas"]:
    if fila["ad_id"] not in conocidos:
        vecino = next((a for a in bib["anuncios"] if a["conjunto"] == fila["conjunto"]), {})
        bib["anuncios"].append({"id": fila["ad_id"], "nombre": fila["nombre"], "conjunto": fila["conjunto"],
                                "campana": fila["campana"], "cuerpo": vecino.get("cuerpo", ""), "formato": vecino.get("formato", ""),
                                "tipo_hook": "", "concepto": "", "hook": "", "alta": fila["fecha"], "estado": "", "notas": ""})
        conocidos.add(fila["ad_id"])
estado_ult = {}
for fila in sorted(hist["filas"], key=lambda x: x["fecha"]):
    if fila.get("estado"):
        estado_ult[fila["ad_id"]] = fila["estado"]
for a in bib["anuncios"]:
    a["estado"] = estado_ult.get(a["id"], a.get("estado", ""))
BIB.write_text(json.dumps(bib, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
HIST.write_text(json.dumps(hist, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")

orden_ang = []
for a in bib["anuncios"]:
    if a["conjunto"] not in orden_ang:
        orden_ang.append(a["conjunto"])
anuncios = sorted(bib["anuncios"], key=lambda a: (orden_ang.index(a["conjunto"]), a["nombre"]))
NA = len(anuncios)
CF = C0 + NA - 1                       # última fila de Creativos
cuerpos = bib["cuerpos"]
NC = len(cuerpos)
conceptos = []
for a in anuncios:
    if a["concepto"] and a["concepto"] not in conceptos:
        conceptos.append(a["concepto"])
color_ang = {g: COLOR_ANGULO[i % len(COLOR_ANGULO)] for i, g in enumerate(orden_ang)}

wb = Workbook()
panel = wb.active
panel.title = "Panel"
mapa = wb.create_sheet("Mapa")
crea = wb.create_sheet("Creativos")
cuer = wb.create_sheet("Cuerpos")
matr = wb.create_sheet("Matriz")
hst = wb.create_sheet("Histórico")
ped = wb.create_sheet("Pedidos")
conf = wb.create_sheet("Config")
guia = wb.create_sheet("Guía")
for w in wb.worksheets:
    w.sheet_view.showGridLines = False

VERED = {"ESCALAR": (VERDE, "FFFFFF"), "MANTENER": (VERDE_SUAVE, VERDE), "PROMETEDOR": (AMBAR_SUAVE, AMBAR),
         "VIGILAR": (AMBAR_SUAVE, AMBAR), "APAGAR": (ROJO_SUAVE, ROJO), "APRENDIENDO": (FONDO_SUB, GRIS)}


def colorear_veredicto(ws, rango):
    for v, (fo, tx) in VERED.items():
        ws.conditional_formatting.add(rango, CellIsRule(operator="equal", formula=[f'"{v}"'], fill=fill(fo),
                                                        font=Font(name=FUENTE, bold=True, color=tx)))


# ---------------------------------------------------------------- Config
ws = conf
for l, w in zip("ABCDEF", (2, 44, 14, 2, 70, 18)):
    ws.column_dimensions[l].width = w
titulo(ws, "Ajustes", "Azul sobre amarillo = editable. Vienen del dashboard de números; cámbialos si allí cambian.")
cabecera(ws, 3, 2, ["Parámetro", "Valor", "", "Qué significa"])
params = [
    (4, "CPA máximo rentable (break-even)", "cpa_be", EUR, "CPA_BE", "Del dashboard (pestaña Números). Por encima de este CPA pierdes dinero."),
    (5, "Beneficio esperado por pedido (antes de Ads)", "ben_pedido", EUR, "BEN_PEDIDO", "Del dashboard: lo que deja de media un pedido confirmado tras IVA, envío y devoluciones."),
    (6, "Recargo sobre el gasto en Ads", "recargo_ads", PCT, "RECARGO", "0% si Meta factura sin IVA o lo deduces."),
    (7, "Fecha de inicio", "fecha_inicio", FECHA, "FECHA_INICIO", "Primer día del histórico."),
    (8, "ESCALAR si el CPA real es ≤ este % del break-even", "umbral_escalar", PCT, "UMB_ESC", "Y al menos 3 compras."),
    (9, "APAGAR sin ventas al gastar × break-even", "apagar_sin_ventas", DEC, "APG_SV", "1,5 = apagar a ~18 € sin una venta."),
    (10, "APAGAR con ventas caras al gastar × break-even", "apagar_con_ventas", DEC, "APG_CV", "Si el CPA real supera el break-even."),
    (11, "Impresiones mínimas para el diagnóstico", "min_impr_diag", ENT, "MIN_IMPR", "Con menos, el diagnóstico dice 'pocos datos'."),
]
for f, etq, k, fmt, n, nota in params:
    celda(ws, f"B{f}", etq)
    v = dt.date.fromisoformat(cfg[k]) if k == "fecha_inicio" else cfg[k]
    entrada(ws, f"C{f}", v, fmt)
    celda(ws, f"E{f}", nota, color=GRIS, size=9, wrap=True)
    celda(ws, f"F{f}", k, color="FFFFFF", size=7)   # clave interna (para recoger cambios)
    nombre(wb, n, f"Config!$C${f}")
celda(ws, "B13", "Listas para los desplegables", bold=True, color=TINTA)
celda(ws, "B14", "Tipos de hook", color=GRIS)
for i, t in enumerate(["Texto", "Visual", "Texto + visual"]):
    celda(ws, f"C{14 + i}", t)
nombre(wb, "LISTA_TIPOS", "Config!$C$14:$C$16")

# ---------------------------------------------------------------- Histórico
ws = hst
cols_h = [("A", "Fecha", 11, FECHA), ("B", "ID anuncio", 19, "@"), ("C", "Anuncio", 20, None), ("D", "Ángulo (conjunto)", 30, None),
          ("E", "Campaña", 30, None), ("F", "Gasto", 9, EUR), ("G", "Impresiones", 10, ENT), ("H", "Clics enlace", 8, ENT),
          ("I", "Visitas web", 8, ENT), ("J", "Añadir carrito", 8, ENT), ("K", "Inicio pago", 8, ENT), ("L", "Compras (píxel)", 8, ENT),
          ("M", "Vídeo 25%", 8, ENT), ("N", "Vídeo 50%", 8, ENT), ("O", "ThruPlays", 8, ENT), ("P", "Tiempo medio (s)", 8, ENT),
          ("Q", "Frecuencia", 8, DEC), ("R", "Cuerpo", 8, None), ("S", "Concepto hook", 14, None), ("T", "Tiempo × impr.", 9, ENT),
          ("U", "CTR", 7, PCT), ("V", "Gancho", 7, PCT), ("W", "CPA píxel", 8, EUR)]
titulo(ws, "Histórico diario por anuncio (Meta Ads)",
       f"Lo carga Claude desde el conector de Meta. Actualizado: {hist.get('actualizado', '')}. {hist.get('nota', '')}")
for l, cab, w, _ in cols_h:
    ws.column_dimensions[l].width = w
cabecera(ws, 4, 1, [c[1] for c in cols_h])
ws.row_dimensions[4].height = 30
filas = sorted(hist["filas"], key=lambda x: (x["fecha"], x["nombre"]))
CLAVES = ["gasto", "impr", "clics", "lpv", "atc", "checkout", "compras", "v25", "v50", "thru", "tmedio", "freq"]
for i, fi in enumerate(filas):
    r = H0 + i
    celda(ws, f"A{r}", dt.date.fromisoformat(fi["fecha"]), fmt=FECHA, borde=BORDE)
    celda(ws, f"B{r}", str(fi["ad_id"]), fmt="@", borde=BORDE, color=GRIS, size=8)
    celda(ws, f"C{r}", fi["nombre"], borde=BORDE)
    celda(ws, f"D{r}", fi["conjunto"], borde=BORDE, size=9)
    celda(ws, f"E{r}", fi["campana"], borde=BORDE, size=8, color=GRIS)
    for letra, k in zip("FGHIJKLMNOPQ", CLAVES):
        celda(ws, f"{letra}{r}", fi.get(k) or 0, fmt=dict((c[0], c[3]) for c in cols_h)[letra], borde=BORDE)
    celda(ws, f"R{r}", f'=IFERROR(INDEX(Creativos!$E${C0}:$E${CF},MATCH($B{r},Creativos!$A${C0}:$A${CF},0)),"")', borde=BORDE, color=GRIS)
    celda(ws, f"S{r}", f'=IFERROR(INDEX(Creativos!$H${C0}:$H${CF},MATCH($B{r},Creativos!$A${C0}:$A${CF},0)),"")', borde=BORDE, color=GRIS)
    celda(ws, f"T{r}", f"=P{r}*G{r}", fmt=ENT, borde=BORDE, color=GRIS)
    celda(ws, f"U{r}", f'=IF(G{r}>0,H{r}/G{r},"")', fmt=PCT, borde=BORDE)
    celda(ws, f"V{r}", f'=IF(G{r}>0,M{r}/G{r},"")', fmt=PCT, borde=BORDE)
    celda(ws, f"W{r}", f'=IF(L{r}>0,F{r}/L{r},"")', fmt=EUR, borde=BORDE)
ws.freeze_panes = f"D{H0}"
ws.auto_filter.ref = f"A4:W{max(H0, H0 + len(filas) - 1)}"

# ---------------------------------------------------------------- Pedidos
ws = ped
titulo(ws, "Pedidos reales por día (Shopify)", "Sirve para corregir las compras de Meta: el píxel no ve todas las ventas.")
for l, w in zip("ABCDEF", (12, 14, 14, 12, 12, 12)):
    ws.column_dimensions[l].width = w
cabecera(ws, 4, 1, ["Fecha", "Pedidos reales (Shopify)", "Compras Meta", "Real ÷ Meta", "Gasto", "CPA real"], entradas=(2,))
fechas = sorted(set(hist["pedidos_reales"]) | {f["fecha"] for f in hist["filas"]})
for i, f in enumerate(fechas):
    r = C0 + i
    celda(ws, f"A{r}", dt.date.fromisoformat(f), fmt=FECHA, borde=BORDE)
    entrada(ws, f"B{r}", hist["pedidos_reales"].get(f), ENT)
    celda(ws, f"C{r}", f"=SUMIFS({HIS}!$L${H0}:$L${HN},{HIS}!$A${H0}:$A${HN},A{r})", fmt=ENT, borde=BORDE)
    celda(ws, f"D{r}", f'=IF(AND(C{r}>0,B{r}<>""),B{r}/C{r},"")', fmt=DEC, borde=BORDE)
    celda(ws, f"E{r}", f"=SUMIFS({HIS}!$F${H0}:$F${HN},{HIS}!$A${H0}:$A${HN},A{r})", fmt=EUR, borde=BORDE)
    celda(ws, f"F{r}", f'=IF(N(B{r})>0,E{r}/B{r},"")', fmt=EUR, borde=BORDE)
PF = C0 + len(fechas) - 1

# ---------------------------------------------------------------- Creativos
ws = crea
titulo(ws, "Creativos: estructura y rendimiento",
       "Una fila por anuncio. Lo amarillo lo editas tú (hook, concepto, cuerpo, notas). Las métricas son del periodo elegido en el Panel.")
cols_c = [("A", "ID anuncio", 19), ("B", "Anuncio", 20), ("C", "Campaña", 22), ("D", "Ángulo (conjunto)", 26),
          ("E", "Cuerpo", 8), ("F", "Formato", 22), ("G", "Tipo de hook", 10), ("H", "Concepto del hook", 16),
          ("I", "Hook (texto, o qué se ve en los 3 primeros segundos)", 44), ("J", "Alta", 10), ("K", "Estado", 9),
          ("L", "Notas / aprendizaje", 30), ("M", "Gasto", 9), ("N", "Impresiones", 10), ("O", "Clics", 7), ("P", "Visitas web", 8),
          ("Q", "Compras Meta", 8), ("R", "Pedidos reales est.", 9), ("S", "CPA real", 9), ("T", "Beneficio est.", 10),
          ("U", "Gancho (ven el 25%)", 9), ("V", "Retención (25%→50%)", 9), ("W", "CTR", 7), ("X", "Llega a la web", 8),
          ("Y", "Compra / clic", 8), ("Z", "Tiempo medio (s)", 8), ("AA", "Veredicto", 13), ("AB", "Diagnóstico", 36), ("AC", "orden", 6)]
for l, cab, w in cols_c:
    ws.column_dimensions[l].width = w
ws.column_dimensions["AC"].hidden = True
cabecera(ws, 4, 1, [c[1] for c in cols_c], entradas=(5, 6, 7, 8, 9, 12))
ws.row_dimensions[4].height = 42
crit = "$A{r}"
for i, a in enumerate(anuncios):
    r = C0 + i
    k = crit.format(r=r)
    banda = color_ang[a["conjunto"]]
    celda(ws, f"A{r}", a["id"], fmt="@", size=8, color=GRIS, borde=BORDE, fondo=banda)
    celda(ws, f"B{r}", a["nombre"], bold=True, borde=BORDE, fondo=banda)
    celda(ws, f"C{r}", a.get("campana", ""), size=8, color=GRIS, borde=BORDE, wrap=True)
    celda(ws, f"D{r}", a["conjunto"], size=9, borde=BORDE, fondo=banda, wrap=True)
    entrada(ws, f"E{r}", a.get("cuerpo", ""))
    entrada(ws, f"F{r}", a.get("formato", ""), wrap=True)
    entrada(ws, f"G{r}", a.get("tipo_hook", ""))
    entrada(ws, f"H{r}", a.get("concepto", ""))
    entrada(ws, f"I{r}", a.get("hook", ""), wrap=True,
            nota="Hook de texto: la frase exacta. Hook visual: describe qué se ve en los 3 primeros segundos." if i == 0 else None)
    celda(ws, f"J{r}", dt.date.fromisoformat(a["alta"]) if a.get("alta") else None, fmt=FECHA, borde=BORDE)
    celda(ws, f"K{r}", {"ACTIVE": "Activo", "PAUSED": "Pausado"}.get(a.get("estado", ""), a.get("estado", "").lower().replace("_", " ")), borde=BORDE, size=9)
    entrada(ws, f"L{r}", a.get("notas", ""), wrap=True)
    f = {
        "M": f"={S('F', 'B', k)}", "N": f"={S('G', 'B', k)}", "O": f"={S('H', 'B', k)}", "P": f"={S('I', 'B', k)}",
        "Q": f"={S('L', 'B', k)}", "R": f"=Q{r}*RATIO", "S": f'=IF(R{r}>0,M{r}/R{r},"")',
        "T": f"=R{r}*BEN_PEDIDO-M{r}*(1+RECARGO)",
        "U": f'=IF(N{r}>0,{S("M", "B", k)}/N{r},"")',
        "V": f'=IF({S("M", "B", k)}>0,{S("N", "B", k)}/{S("M", "B", k)},"")',
        "W": f'=IF(N{r}>0,O{r}/N{r},"")', "X": f'=IF(O{r}>0,P{r}/O{r},"")', "Y": f'=IF(O{r}>0,Q{r}/O{r},"")',
        "Z": f'=IF(N{r}>0,{S("T", "B", k)}/N{r},"")',
        "AA": (f'=IF(M{r}=0,"SIN GASTO",IF(Q{r}=0,IF(M{r}>=APG_SV*CPA_BE,"APAGAR",IF(M{r}>=0.5*CPA_BE,"VIGILAR","APRENDIENDO")),'
               f'IF(AND(S{r}<=UMB_ESC*CPA_BE,Q{r}>=3),"ESCALAR",IF(AND(S{r}<=CPA_BE,Q{r}>=3),"MANTENER",'
               f'IF(S{r}<=CPA_BE,"PROMETEDOR",IF(M{r}>=APG_CV*CPA_BE,"APAGAR","VIGILAR"))))))'),
        "AB": (f'=IF(N{r}<MIN_IMPR,"Pocos datos todavía",IF(AND(ISNUMBER(U{r}),U{r}<0.7*MEDIA_GANCHO),"Gancho flojo: cambia los 3 primeros segundos",'
               f'IF(AND(ISNUMBER(W{r}),W{r}<0.7*MEDIA_CTR),"Enganchan pero no hacen clic: refuerza oferta y llamada a la acción",'
               f'IF(AND(O{r}>=20,ISNUMBER(Y{r}),Y{r}<0.6*MEDIA_CVR),"Hacen clic pero no compran: revisa ángulo y página",'
               f'IF(AND(ISNUMBER(U{r}),U{r}>=1.3*MEDIA_GANCHO),"Gancho por encima de la media: reutiliza este inicio","En la media de la campaña")))))'),
        "AC": f"=IF(M{r}>0,T{r}+ROW()/1000000,-1000000+ROW())",
    }
    fmts = {"M": EUR, "N": ENT, "O": ENT, "P": ENT, "Q": ENT, "R": '0.0;-0.0;"-"', "S": EUR, "T": EUR, "U": PCT, "V": PCT,
            "W": PCT, "X": PCT, "Y": PCT, "Z": '0.0;-0.0;"-"', "AA": None, "AB": None, "AC": "0.00"}
    for letra, formula in f.items():
        celda(ws, f"{letra}{r}", formula, fmt=fmts[letra], borde=BORDE, bold=letra in ("S", "T", "AA"),
              wrap=letra == "AB", size=9 if letra == "AB" else 10, align="center" if letra == "AA" else None)
    ws.row_dimensions[r].height = 46
dv_c = DataValidation(type="list", formula1=f"=Cuerpos!$A${C0}:$A${C0 + NC + 9}", allow_blank=True)
dv_t = DataValidation(type="list", formula1="=LISTA_TIPOS", allow_blank=True)
ws.add_data_validation(dv_c)
ws.add_data_validation(dv_t)
dv_c.add(f"E{C0}:E{CF + 20}")
dv_t.add(f"G{C0}:G{CF + 20}")
colorear_veredicto(ws, f"AA{C0}:AA{CF}")
ws.conditional_formatting.add(f"T{C0}:T{CF}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=FUENTE, bold=True, color=VERDE)))
ws.conditional_formatting.add(f"S{C0}:S{CF}", FormulaRule(formula=[f"AND(ISNUMBER(S{C0}),S{C0}>CPA_BE)"], font=Font(name=FUENTE, bold=True, color=ROJO)))
for letra, media in (("U", "MEDIA_GANCHO"), ("W", "MEDIA_CTR"), ("X", "MEDIA_WEB"), ("Y", "MEDIA_CVR")):
    ws.conditional_formatting.add(f"{letra}{C0}:{letra}{CF}", FormulaRule(
        formula=[f"AND(ISNUMBER({letra}{C0}),$N{C0}>=300,{letra}{C0}>=1.3*{media})"], font=Font(name=FUENTE, bold=True, color=VERDE)))
    ws.conditional_formatting.add(f"{letra}{C0}:{letra}{CF}", FormulaRule(
        formula=[f"AND(ISNUMBER({letra}{C0}),$N{C0}>=300,{letra}{C0}<0.7*{media})"], font=Font(name=FUENTE, bold=True, color=ROJO)))
ws.freeze_panes = f"C{C0}"

# ---------------------------------------------------------------- Cuerpos
ws = cuer
titulo(ws, "Cuerpos: el guion después del hook, dividido por partes",
       "Edita cada parte en amarillo; el texto completo, la duración y el rendimiento se calculan solos. Para un cuerpo nuevo, escribe un ID nuevo (C4, C5…) en la primera fila libre.")
cols_b = [("A", "ID", 6), ("B", "Nombre", 18), ("C", "Ángulo", 18), ("D", "Formato", 16), ("E", "1 · Apertura / problema", 30),
          ("F", "2 · Giro / solución", 26), ("G", "3 · Mecanismo (por qué funciona)", 30), ("H", "4 · Prueba / filtro / garantía", 28),
          ("I", "5 · Oferta", 22), ("J", "6 · Llamada a la acción", 22), ("K", "Texto completo", 50), ("L", "Palabras", 8),
          ("M", "Duración aprox. (s)", 9), ("N", "Precio que menciona", 11), ("O", "Notas", 24), ("P", "Anuncios", 8),
          ("Q", "Gasto", 9), ("R", "Compras Meta", 8), ("S", "CPA real", 9), ("T", "Beneficio est.", 10)]
for l, cab, w in cols_b:
    ws.column_dimensions[l].width = w
cabecera(ws, 4, 1, [c[1] for c in cols_b], entradas=(1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 14, 15))
ws.row_dimensions[4].height = 40
for i in range(NC + 10):
    r = C0 + i
    c = cuerpos[i] if i < NC else {}
    entrada(ws, f"A{r}", c.get("id"))
    for letra, k in zip("BCDEFGHIJ", ("nombre", "angulo", "formato", "apertura", "solucion", "mecanismo", "prueba", "oferta", "cta")):
        entrada(ws, f"{letra}{r}", c.get(k) or None, wrap=True)
    entrada(ws, f"N{r}", c.get("precio") or None, wrap=True)
    entrada(ws, f"O{r}", c.get("notas") or None, wrap=True)
    celda(ws, f"K{r}", f'=TRIM(E{r}&" "&F{r}&" "&G{r}&" "&H{r}&" "&I{r}&" "&J{r})', wrap=True, size=9, borde=BORDE, valign="top")
    celda(ws, f"L{r}", f'=IF(K{r}="","",LEN(K{r})-LEN(SUBSTITUTE(K{r}," ",""))+1)', fmt=ENT, borde=BORDE)
    celda(ws, f"M{r}", f'=IF(L{r}="","",ROUND(L{r}/2.6,0))', fmt=ENT, borde=BORDE)
    celda(ws, f"P{r}", f'=IF(A{r}="","",COUNTIF(Creativos!$E${C0}:$E${CF},A{r}))', fmt=ENT, borde=BORDE)
    celda(ws, f"Q{r}", f'=IF(A{r}="","",{S("F", "R", f"$A{r}")})', fmt=EUR, borde=BORDE)
    celda(ws, f"R{r}", f'=IF(A{r}="","",{S("L", "R", f"$A{r}")})', fmt=ENT, borde=BORDE)
    celda(ws, f"S{r}", f'=IF(N(R{r})*RATIO>0,Q{r}/(R{r}*RATIO),"")', fmt=EUR, borde=BORDE, bold=True)
    celda(ws, f"T{r}", f'=IF(A{r}="","",R{r}*RATIO*BEN_PEDIDO-Q{r}*(1+RECARGO))', fmt=EUR, borde=BORDE, bold=True)
    ws.row_dimensions[r].height = 150 if i < NC else 30
ws.conditional_formatting.add(f"T{C0}:T{C0 + NC + 9}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=FUENTE, bold=True, color=VERDE)))
ws.conditional_formatting.add(f"T{C0}:T{C0 + NC + 9}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FUENTE, bold=True, color=ROJO)))
ws.freeze_panes = f"C{C0}"

# ---------------------------------------------------------------- Panel
ws = panel
for i, w in enumerate([2] + [11.5] * 11 + [15, 2]):
    ws.column_dimensions[col(i + 1)].width = w
titulo(ws, "CeraLux · Creativos", "Qué anuncio, qué hook y qué cuerpo te están dando dinero. Datos de Meta corregidos con tus pedidos reales.")
celda(ws, "B3", "Desde", bold=True, color=TINTA, align="right")
entrada(ws, "C3", "=FECHA_INICIO", FECHA, nota="Escribe una fecha para fijar el periodo.")
ws.merge_cells("C3:D3")
celda(ws, "E3", "Hasta", bold=True, color=TINTA, align="right")
entrada(ws, "F3", "=TODAY()", FECHA)
ws.merge_cells("F3:G3")
celda(ws, "H3", "Días", bold=True, color=TINTA, align="right")
celda(ws, "I3", "=MAX(1,F3-C3+1)", fmt=ENT, bold=True)
celda(ws, "J3", f"Actualizado: {hist.get('actualizado', '')[:16].replace('T', ' ')}", color=GRIS, size=9, italic=True)
nombre(wb, "DESDE", "Panel!$C$3")
nombre(wb, "HASTA", "Panel!$F$3")

seccion(ws, "B5:M5", "RESUMEN DEL PERIODO")
PEDIDOS_P = f"SUMIFS(Pedidos!$B${C0}:$B${PF},Pedidos!$A${C0}:$A${PF},\">=\"&DESDE,Pedidos!$A${C0}:$A${PF},\"<=\"&HASTA)"
TOT = lambda letra: (f"SUMIFS({HIS}!${letra}${H0}:${letra}${HN},{HIS}!$A${H0}:$A${HN},\">=\"&DESDE,"  # noqa: E731
                     f"{HIS}!$A${H0}:$A${HN},\"<=\"&HASTA)")
TARJ = [
    [("Gasto", f"={TOT('F')}", EUR), ("Compras (Meta)", f"={TOT('L')}", ENT),
     ("Pedidos reales", f"=IF({PEDIDOS_P}>0,{PEDIDOS_P},D7)", ENT), ("Pedidos reales ÷ Meta", '=IF(D7>0,F7/D7,1)', DEC),
     ("CPA real", '=IF(F7>0,B7/F7,"")', EUR), ("Beneficio estimado", "=F7*BEN_PEDIDO-B7*(1+RECARGO)", EUR)],
    [("CPA máximo rentable", "=CPA_BE", EUR), ("Gancho medio (25%)", f"=IFERROR({TOT('M')}/{TOT('G')},0)", PCT),
     ("CTR medio", f"=IFERROR({TOT('H')}/{TOT('G')},0)", PCT), ("Llega a la web", f"=IFERROR({TOT('I')}/{TOT('H')},0)", PCT),
     ("Compra / clic", f"=IFERROR({TOT('L')}/{TOT('H')},0)", PCT), ("Anuncios con gasto", f'=COUNTIF(Creativos!$M${C0}:$M${CF},">0")', ENT)],
]
for fila_t, grupo in zip((6, 9), TARJ):
    for j, (etq, formula, fmt) in enumerate(grupo):
        c1, c2 = col(2 + 2 * j), col(3 + 2 * j)
        ws.merge_cells(f"{c1}{fila_t}:{c2}{fila_t}")
        ws.merge_cells(f"{c1}{fila_t + 1}:{c2}{fila_t + 1}")
        celda(ws, f"{c1}{fila_t}", etq.upper(), bold=True, color=GRIS, size=8, fondo=FONDO_CARD, align="center")
        celda(ws, f"{c1}{fila_t + 1}", formula, bold=True, color=TINTA, size=15, fmt=fmt, fondo=FONDO_CARD, align="center")
    ws.row_dimensions[fila_t + 1].height = 30
nombre(wb, "RATIO", "Panel!$H$7")
nombre(wb, "MEDIA_GANCHO", "Panel!$D$10")
nombre(wb, "MEDIA_CTR", "Panel!$F$10")
nombre(wb, "MEDIA_WEB", "Panel!$H$10")
nombre(wb, "MEDIA_CVR", "Panel!$J$10")
ws.conditional_formatting.add("J7", FormulaRule(formula=["AND(ISNUMBER(J7),J7>CPA_BE)"], font=Font(name=FUENTE, bold=True, color=ROJO, size=15)))
ws.conditional_formatting.add("L7", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=FUENTE, bold=True, color=VERDE, size=15)))

seccion(ws, "B12:M12", "QUÉ HACER CON TUS ANUNCIOS")
for j, v in enumerate(["ESCALAR", "MANTENER", "PROMETEDOR", "VIGILAR", "APAGAR", "APRENDIENDO"]):
    c1, c2 = col(2 + 2 * j), col(3 + 2 * j)
    ws.merge_cells(f"{c1}13:{c2}13")
    ws.merge_cells(f"{c1}14:{c2}14")
    fo, tx = VERED[v]
    celda(ws, f"{c1}13", v, bold=True, color=tx, size=9, fondo=fo, align="center")
    celda(ws, f"{c1}14", f'=COUNTIF(Creativos!$AA${C0}:$AA${CF},"{v}")', bold=True, color=TINTA, size=14, fmt=ENT, fondo=FONDO_CARD, align="center")

seccion(ws, "B16:M16", "TOP 5 CREATIVOS POR BENEFICIO ESTIMADO")
cabecera(ws, 17, 2, ["#", "Anuncio", "", "Ángulo", "", "Hook", "", "", "Gasto", "CPA real", "Beneficio", "Veredicto"])
for a, b in (("C17", "D17"), ("E17", "F17"), ("G17", "I17")):
    ws.merge_cells(f"{a}:{b}")
CRA = lambda letra: f"Creativos!${letra}${C0}:${letra}${CF}"  # noqa: E731
for k in range(1, 6):
    r = 17 + k
    pos = f"MATCH(LARGE({CRA('AC')},{k}),{CRA('AC')},0)"
    celda(ws, f"B{r}", k, align="center", bold=True, borde=BORDE)
    ws.merge_cells(f"C{r}:D{r}")
    celda(ws, f"C{r}", f"=IFERROR(INDEX({CRA('B')},{pos}),\"\")", bold=True, borde=BORDE)
    ws.merge_cells(f"E{r}:F{r}")
    celda(ws, f"E{r}", f"=IFERROR(INDEX({CRA('D')},{pos}),\"\")", size=9, borde=BORDE, wrap=True)
    ws.merge_cells(f"G{r}:I{r}")
    celda(ws, f"G{r}", f'=IFERROR(IF(INDEX({CRA("I")},{pos})="","[visual] "&INDEX({CRA("H")},{pos}),INDEX({CRA("I")},{pos})),"")',
          size=9, borde=BORDE, wrap=True)
    celda(ws, f"J{r}", f"=IFERROR(INDEX({CRA('M')},{pos}),\"\")", fmt=EUR, borde=BORDE)
    celda(ws, f"K{r}", f"=IFERROR(INDEX({CRA('S')},{pos}),\"\")", fmt=EUR, borde=BORDE, bold=True)
    celda(ws, f"L{r}", f"=IFERROR(INDEX({CRA('T')},{pos}),\"\")", fmt=EUR, borde=BORDE, bold=True)
    celda(ws, f"M{r}", f"=IFERROR(INDEX({CRA('AA')},{pos}),\"\")", borde=BORDE, align="center")
    ws.row_dimensions[r].height = 30
colorear_veredicto(ws, "M18:M22")


def tabla_resumen(fila, titulo_s, etiquetas, crit_col, etq_col, nombres=None):
    seccion(ws, f"B{fila}:M{fila}", titulo_s)
    cabecera(ws, fila + 1, 2, [etq_col, "", "", "", "Anuncios", "Gasto", "Compras Meta", "Pedidos est.", "CPA real", "Beneficio est.", "Gancho", "CTR"])
    ws.merge_cells(f"B{fila + 1}:E{fila + 1}")
    for i, e in enumerate(etiquetas):
        r = fila + 2 + i
        ws.merge_cells(f"B{r}:E{r}")
        if nombres:   # "C1 · Rabia · el bote habla": se filtra por lo que va antes del primer " · "
            celda(ws, f"B{r}", f"{e} · {nombres[e]}" if nombres.get(e) else e, borde=BORDE, bold=True)
            crit = f'LEFT($B{r},FIND(" · ",$B{r}&" · ")-1)'
        else:
            celda(ws, f"B{r}", e, borde=BORDE, bold=True)
            crit = f"$B{r}"
        cuenta_col = {"D": "D", "R": "E", "S": "H"}[crit_col]
        celda(ws, f"F{r}", f"=COUNTIF({CRA(cuenta_col)},{crit})", fmt=ENT, borde=BORDE)
        celda(ws, f"G{r}", f"={S('F', crit_col, crit)}", fmt=EUR, borde=BORDE)
        celda(ws, f"H{r}", f"={S('L', crit_col, crit)}", fmt=ENT, borde=BORDE)
        celda(ws, f"I{r}", f"=H{r}*RATIO", fmt='0.0;-0.0;"-"', borde=BORDE)
        celda(ws, f"J{r}", f'=IF(I{r}>0,G{r}/I{r},"")', fmt=EUR, borde=BORDE, bold=True)
        celda(ws, f"K{r}", f"=I{r}*BEN_PEDIDO-G{r}*(1+RECARGO)", fmt=EUR, borde=BORDE, bold=True)
        celda(ws, f"L{r}", f'=IFERROR({S("M", crit_col, crit)}/{S("G", crit_col, crit)},"")', fmt=PCT, borde=BORDE)
        celda(ws, f"M{r}", f'=IFERROR({S("H", crit_col, crit)}/{S("G", crit_col, crit)},"")', fmt=PCT, borde=BORDE)
    fin = fila + 1 + len(etiquetas)
    ws.conditional_formatting.add(f"K{fila + 2}:K{fin}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=FUENTE, bold=True, color=VERDE)))
    ws.conditional_formatting.add(f"K{fila + 2}:K{fin}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FUENTE, bold=True, color=ROJO)))
    return fin + 2


fila = 24
fila = tabla_resumen(fila, "POR ÁNGULO (conjunto de anuncios)", orden_ang, "D", "Ángulo")
fila = tabla_resumen(fila, "POR CUERPO (guion después del hook)", [c["id"] for c in cuerpos], "R", "Cuerpo",
                    {c["id"]: c.get("nombre", "") for c in cuerpos})
fila = tabla_resumen(fila, "POR CONCEPTO DE HOOK", conceptos, "S", "Concepto del hook")

graf = BarChart()
graf.type = "bar"
graf.style = 10
graf.title = "Beneficio estimado por anuncio (€)"
graf.add_data(Reference(crea, min_col=20, min_row=4, max_row=CF), titles_from_data=True)
graf.set_categories(Reference(crea, min_col=2, min_row=C0, max_row=CF))
graf.series[0].graphicalProperties.solidFill = TINTA
graf.legend = None
graf.gapWidth = 50
graf.x_axis.tickLblPos = "low"   # nombres a la izquierda aunque el beneficio sea negativo
graf.x_axis.delete = False
graf.y_axis.delete = False
graf.y_axis.number_format = '#,##0 "€"'
graf.height, graf.width = max(7, 0.65 * NA), 31
ws.add_chart(graf, f"B{fila}")
for r in range(1, fila + 40):     # alturas fijas: si no, LibreOffice/Excel recolocan el gráfico encima de la tabla
    if not ws.row_dimensions[r].height:
        ws.row_dimensions[r].height = 24 if r == 1 else 16
ws.freeze_panes = "A4"

# ---------------------------------------------------------------- Mapa
ws = mapa
for l, w in zip("ABCDEFGHIJ", (2, 58, 20, 46, 10, 9, 10, 12, 15, 2)):
    ws.column_dimensions[l].width = w
titulo(ws, "Mapa de creativos: ángulo → cuerpo → hooks",
       "Tu pizarra de estructura con los números al lado. Se rellena sola desde Creativos y Cuerpos (periodo del Panel).")
fila_crea = {a["id"]: C0 + i for i, a in enumerate(anuncios)}
texto_cuerpo = {c["id"]: " ".join(x for x in (c.get(k, "") for k in ("apertura", "solucion", "mecanismo", "prueba", "oferta", "cta")) if x) for c in cuerpos}
nombre_cuerpo = {c["id"]: c.get("nombre", "") for c in cuerpos}
grupos = []
for a in anuncios:
    clave = (a["conjunto"], a.get("cuerpo", ""))
    if not grupos or grupos[-1][0] != clave:
        grupos.append((clave, []))
    grupos[-1][1].append(a)
r = 4
for (conjunto, cid), ads in grupos:
    banda = color_ang[conjunto]
    ws.merge_cells(f"B{r}:I{r}")
    celda(ws, f"B{r}", f"{conjunto}   ·   Cuerpo {cid or '—'}: {nombre_cuerpo.get(cid, '')}", bold=True, color=TINTA, size=12, fondo=banda)
    ws.row_dimensions[r].height = 24
    r += 1
    cabecera(ws, r, 2, ["CUERPO (lo mismo para todos sus hooks)", "ANUNCIO", "HOOK", "Gasto", "Compras", "CPA real", "Beneficio est.", "Veredicto"])
    r += 1
    r0 = r
    for a in ads:
        rc = fila_crea[a["id"]]
        celda(ws, f"C{r}", f"=Creativos!B{rc}", bold=True, borde=BORDE, fondo=banda)
        celda(ws, f"D{r}", f'=IF(Creativos!I{rc}="","[visual] "&Creativos!H{rc},Creativos!I{rc})', borde=BORDE, wrap=True, size=9)
        celda(ws, f"E{r}", f"=Creativos!M{rc}", fmt=EUR, borde=BORDE)
        celda(ws, f"F{r}", f"=Creativos!Q{rc}", fmt=ENT, borde=BORDE)
        celda(ws, f"G{r}", f"=Creativos!S{rc}", fmt=EUR, borde=BORDE, bold=True)
        celda(ws, f"H{r}", f"=Creativos!T{rc}", fmt=EUR, borde=BORDE, bold=True)
        celda(ws, f"I{r}", f"=Creativos!AA{rc}", borde=BORDE, align="center")
        r += 1
    r1 = r - 1
    ws.merge_cells(f"B{r0}:B{r1}")
    celda(ws, f"B{r0}", f'=IFERROR(INDEX(Cuerpos!$K${C0}:$K${C0 + NC + 9},MATCH("{cid}",Cuerpos!$A${C0}:$A${C0 + NC + 9},0)),"")',
          wrap=True, size=9, borde=BORDE, valign="top", fondo="FFFDF5")
    lineas = len(texto_cuerpo.get(cid, "")) / 75 + 1
    alto = max(34, lineas * 12.5 / len(ads))
    for rr in range(r0, r1 + 1):
        ws.row_dimensions[rr].height = alto
    celda(ws, f"C{r}", "Total", bold=True, color=TINTA, fondo=FONDO_CARD)
    celda(ws, f"E{r}", f"=SUM(E{r0}:E{r1})", fmt=EUR, bold=True, fondo=FONDO_CARD)
    celda(ws, f"F{r}", f"=SUM(F{r0}:F{r1})", fmt=ENT, bold=True, fondo=FONDO_CARD)
    celda(ws, f"G{r}", f'=IF(F{r}*RATIO>0,E{r}/(F{r}*RATIO),"")', fmt=EUR, bold=True, fondo=FONDO_CARD)
    celda(ws, f"H{r}", f"=SUM(H{r0}:H{r1})", fmt=EUR, bold=True, fondo=FONDO_CARD)
    colorear_veredicto(ws, f"I{r0}:I{r1}")
    ws.conditional_formatting.add(f"H{r0}:H{r}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FUENTE, bold=True, color=ROJO)))
    ws.conditional_formatting.add(f"H{r0}:H{r}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=FUENTE, bold=True, color=VERDE)))
    r += 2

# ---------------------------------------------------------------- Matriz hook × cuerpo
ws = matr
titulo(ws, "Matriz hook × cuerpo: CPA real de cada combinación",
       "Cada celda es un concepto de hook montado sobre un cuerpo. '—' = combinación sin probar: ahí están tus próximos tests.")
ws.column_dimensions["A"].width = 2
ws.column_dimensions["B"].width = 22
for j in range(NC):
    ws.column_dimensions[col(3 + j)].width = 18
cabecera(ws, 4, 2, ["Concepto del hook"] + [f"{c['id']} · {c.get('nombre', '')}" for c in cuerpos])
ws.row_dimensions[4].height = 36
for i, cpt in enumerate(conceptos):
    r = 5 + i
    celda(ws, f"B{r}", cpt, bold=True, borde=BORDE)
    for j, c in enumerate(cuerpos):
        L = col(3 + j)
        gasto = (f"SUMIFS({HIS}!$F${H0}:$F${HN},{HIS}!$S${H0}:$S${HN},$B{r},{HIS}!$R${H0}:$R${HN},\"{c['id']}\","
                 f"{HIS}!$A${H0}:$A${HN},\">=\"&DESDE,{HIS}!$A${H0}:$A${HN},\"<=\"&HASTA)")
        compras = gasto.replace("$F$", "$L$")
        celda(ws, f"{L}{r}", f'=IF({gasto}=0,"—",IF({compras}=0,"sin ventas",{gasto}/({compras}*RATIO)))', fmt=EUR,
              borde=BORDE, align="center", bold=True)
    ws.row_dimensions[r].height = 26
fin = 4 + len(conceptos)
rng = f"C5:{col(2 + NC)}{fin}"
ws.conditional_formatting.add(rng, FormulaRule(formula=["AND(ISNUMBER(C5),C5<=UMB_ESC*CPA_BE)"], fill=fill(VERDE_SUAVE), font=Font(name=FUENTE, bold=True, color=VERDE)))
ws.conditional_formatting.add(rng, FormulaRule(formula=["AND(ISNUMBER(C5),C5>UMB_ESC*CPA_BE,C5<=CPA_BE)"], fill=fill(AMBAR_SUAVE), font=Font(name=FUENTE, bold=True, color=AMBAR)))
ws.conditional_formatting.add(rng, FormulaRule(formula=['OR(AND(ISNUMBER(C5),C5>CPA_BE),C5="sin ventas")'], fill=fill(ROJO_SUAVE), font=Font(name=FUENTE, bold=True, color=ROJO)))
celda(ws, f"B{fin + 2}", "Verde: CPA real ≤ 75% del break-even (escalar). Ámbar: rentable justo. Rojo: pierde dinero o sin ventas.", color=GRIS, size=9, italic=True)

# ---------------------------------------------------------------- Guía
ws = guia
ws.column_dimensions["A"].width = 2
ws.column_dimensions["B"].width = 120
celda(ws, "B1", "Cómo usar este libro", bold=True, color=TINTA, size=16)
lineas = [
    ("h", "Hojas"),
    ("t", "Panel: elige Desde / Hasta y mira qué anuncios, ángulos, cuerpos y hooks dan dinero. El Top 5 y los totales salen solos."),
    ("t", "Mapa: tu pizarra (ángulo → cuerpo → hooks) con gasto, CPA real, beneficio y veredicto al lado de cada hook."),
    ("t", "Creativos: la tabla maestra. Una fila por anuncio. Rellena en amarillo el hook, el concepto, el cuerpo que usa y tus notas."),
    ("t", "Cuerpos: cada guion dividido en 6 partes (apertura, giro, mecanismo, prueba, oferta, llamada a la acción). Cambia una parte y "
          "el texto completo y la duración se recalculan."),
    ("t", "Matriz: qué combinación de hook + cuerpo funciona. Las celdas con '—' son combinaciones que aún no has probado."),
    ("t", "Histórico: los datos diarios de Meta por anuncio. Pedidos: tus pedidos reales de Shopify para corregir lo que Meta no ve."),
    ("h", "Cómo se actualiza"),
    ("t", "Dile a Claude \"actualiza el Excel de creativos\". Trae de Meta el histórico de cada anuncio y de Shopify los pedidos, "
          "regenera el libro y conserva todo lo que hayas escrito a mano (hooks, cuerpos, notas, ajustes)."),
    ("t", "Los anuncios nuevos aparecen solos en Creativos: solo tienes que escribir su hook, su concepto y qué cuerpo usan."),
    ("h", "Cómo se calcula"),
    ("t", "Pedidos reales est. = compras de Meta × (pedidos reales de Shopify ÷ compras de Meta) en el periodo."),
    ("t", "CPA real = gasto ÷ pedidos reales est. Beneficio est. = pedidos reales est. × beneficio por pedido − gasto."),
    ("t", "Gancho = % de impresiones que ven al menos el 25% del vídeo. Retención = de esos, cuántos llegan al 50%."),
    ("t", "Veredicto (igual que en el dashboard): ESCALAR (≥3 compras y CPA ≤ 75% del break-even) · MANTENER (≥3 y ≤ break-even) · "
          "PROMETEDOR (<3 compras y ≤ break-even) · VIGILAR · APAGAR (sin ventas al gastar 1,5 × break-even, o caro al gastar 2 ×) · APRENDIENDO."),
    ("h", "Para sacarle partido"),
    ("t", "Nombra los anuncios ÁNGULO_CONCEPTO_HOOK_VERSIÓN (ej.: BROAD_LLAVES_H2_V3). Así cada variación se agrupa sola por concepto."),
    ("t", "Testea hooks nuevos sobre un cuerpo ganador (mismo cuerpo, distinto inicio): así sabes si lo que gana es el hook o el cuerpo."),
    ("t", "Colores: azul sobre amarillo = lo escribes tú. Negro = fórmula."),
]
fila = 3
for tipo, texto in lineas:
    if tipo == "h":
        fila += 1
        celda(ws, f"B{fila}", texto, bold=True, color=TINTA, size=11)
    else:
        celda(ws, f"B{fila}", texto, wrap=True)
    fila += 1

wb.active = 0
wb.save(SALIDA)
print(f"Guardado: {SALIDA} · {NA} anuncios · {NC} cuerpos · {len(filas)} filas de histórico")
